"""
Reporte de inscripción con columnas 'Identificación' y 'DOCUMENTO' (OK/NO): las dos se reconocen
como identificación. Antes la hoja entera fallaba y la planilla quedaba vacía (todos "Falta en Excel").
"""
from openpyxl import Workbook

from app.services.comparacion_service import comparacion_service
from app.services.excel_lookup_service import excel_lookup_service

FILAS = [
    ("CC - 1006000111", "ALEJANDRO VEGA MORA", "Seleccionado", "OK", "SI"),
    ("CC - 1119000222", "JOSE DAVID LOPEZ RUIZ", "No Seleccionado", "NO", "NO"),
    ("CC - 1010000333", "CAMILO ALFONSO SOTO VARGAS", "Seleccionado", "OK", "SI"),
]


def _reporte(ruta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inscritos Primera Opción"
    ws["A3"], ws["B3"] = "Ficha", "3100003"
    ws["A4"], ws["B4"] = "Programa de Formación", "ESTRUCTURAS EN CONCRETO"
    for j, titulo in enumerate(("Identificación", "Nombre", "Estado", "DOCUMENTO", "FIRMO")):
        ws.cell(row=6, column=1 + j, value=titulo)
    for i, fila in enumerate(FILAS):
        for j, valor in enumerate(fila):
            ws.cell(row=7 + i, column=1 + j, value=valor)
    wb.create_sheet("Inscritos Segunda Opción")["A6"] = "Identificación"
    wb.save(ruta)
    return str(ruta)


def test_lookup_toma_todas_las_personas(tmp_path):
    lookup = excel_lookup_service.cargar_lookup(_reporte(tmp_path / "reporte.xlsx"))
    assert set(lookup) == {"1006000111", "1119000222", "1010000333"}
    assert lookup["1119000222"]["nombre_completo"] == "JOSE DAVID LOPEZ RUIZ"


def test_comparacion_toma_todas_las_personas(tmp_path):
    df = comparacion_service.cargar_excel(_reporte(tmp_path / "reporte.xlsx"))
    assert sorted(df["numero_identificacion"]) == ["1006000111", "1010000333", "1119000222"]
