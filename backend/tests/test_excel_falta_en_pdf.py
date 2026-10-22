"""
Al subir PDF + Excel, quien está en el Excel pero no en el PDF queda en la ficha como "Falta en PDF".
- Funciona con el reporte de inscripción (columnas Identificación y DOCUMENTO).
- Al agregar otra parte de la ficha, las personas de la parte anterior no se marcan como faltantes.
- Una persona que ya está en otra ficha del usuario no se mueve de allí.
"""
import uuid

from openpyxl import Workbook

from app.models.persona import Persona
from tests.test_documentos_soporte import CEDULA_AMARILLA, CEDULA_DIGITAL
from tests.test_versiones_ficha import ocr_simulado  # noqa: F401 (fixture)
from tests.utilidades_pg import cabecera, crear_usuario, pdf_de_prueba, requiere_postgres

ADRIANA, ANYI, JOSE = "1006459309", "1117501209", "1119580110"


def _reporte(ruta, filas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inscritos Primera Opción"
    ws["A3"], ws["B3"] = "Código Ficha", "3640522"
    for j, titulo in enumerate(("Identificación", "Nombre", "Estado", "DOCUMENTO", "FIRMO")):
        ws.cell(row=6, column=1 + j, value=titulo)
    for i, fila in enumerate(filas):
        for j, valor in enumerate(fila):
            ws.cell(row=7 + i, column=1 + j, value=valor)
    wb.create_sheet("Inscritos Segunda Opción")["A6"] = "Identificación"
    wb.save(ruta)
    return ruta


REPORTE = [
    (f"CC - {ADRIANA}", "ADRIANA ALEJANDRA ZAPATA CUELLAR", "Seleccionado", "OK", "SI"),
    (f"CC - {ANYI}", "ANYI YULIANA BALANTA URIBE", "Seleccionado", "OK", "SI"),
    (f"CC - {JOSE}", "JOSE RUBIEL LOAIZA ICO", "No Seleccionado", "NO", "NO"),
]


def _subir(api, usuario, tmp_path, ocr_simulado, nombre, textos, excel=None, acciones=None):
    import json
    ocr_simulado.clear()
    ocr_simulado.update({i + 1: t for i, t in enumerate(textos)})
    pdf = pdf_de_prueba(tmp_path / f"{uuid.uuid4()}.pdf", paginas=len(textos))
    archivos = [("files", (nombre, open(pdf, "rb").read(), "application/pdf"))]
    if excel:
        archivos.append(("excel", ("reporte.xlsx", open(excel, "rb").read(),
                                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")))
    datos = {"acciones": json.dumps(acciones)} if acciones else None
    r = api.post("/api/documentos/upload", headers=usuario, files=archivos, data=datos)
    assert r.status_code == 202, r.text
    return r.json()["documentos"][0]


def _personas(bd):
    db = bd.SessionLocal()
    res = {p.numero_identificacion: (p.documento_id, p.motor_ocr, (p.detalles_campos or {}).get("en_pdf"), p.nombre_completo)
           for p in db.query(Persona).all()}
    db.close()
    return res


@requiere_postgres
def test_quien_no_esta_en_el_pdf_queda_como_falta_en_pdf(bd, api, tmp_path, ocr_simulado):
    db = bd.SessionLocal()
    ana = cabecera(crear_usuario(db, "ana@example.com"))
    db.close()
    excel = _reporte(tmp_path / "reporte.xlsx", REPORTE)

    ficha = _subir(api, ana, tmp_path, ocr_simulado, "3640522.pdf", [CEDULA_AMARILLA, CEDULA_DIGITAL], excel=excel)
    personas = _personas(bd)

    assert set(personas) == {ADRIANA, ANYI, JOSE}
    documento_id, motor, en_pdf, nombre = personas[JOSE]
    assert documento_id == uuid.UUID(ficha["id"])
    assert motor == "excel" and en_pdf is False and nombre == "JOSE RUBIEL LOAIZA ICO"
    assert personas[ADRIANA][1] != "excel" and personas[ANYI][1] != "excel"

    r = api.get("/api/personas", headers=ana, params={"documento_id": ficha["id"], "limite": 100})
    assert r.status_code == 200
    items = r.json()["items"] if isinstance(r.json(), dict) else r.json()
    assert len(items) == 3


@requiere_postgres
def test_al_agregar_otra_parte_no_se_marcan_como_faltantes_los_de_la_parte_anterior(bd, api, tmp_path, ocr_simulado):
    db = bd.SessionLocal()
    ana = cabecera(crear_usuario(db, "ana@example.com"))
    db.close()
    excel = _reporte(tmp_path / "reporte.xlsx", REPORTE)

    parte1 = _subir(api, ana, tmp_path, ocr_simulado, "3640522 parte 1.pdf", [CEDULA_AMARILLA], excel=excel)
    _subir(api, ana, tmp_path, ocr_simulado, "3640522 parte 2.pdf", [CEDULA_DIGITAL], excel=excel,
           acciones=[{"archivo": "3640522 parte 2.pdf", "accion": "agregar", "documento_id": parte1["id"]}])
    personas = _personas(bd)

    assert set(personas) == {ADRIANA, ANYI, JOSE}
    assert all(p[0] == uuid.UUID(parte1["id"]) for p in personas.values())
    assert personas[ADRIANA][1] != "excel"                              # sigue con sus datos del PDF
    assert personas[ANYI][1] != "excel"
    assert personas[JOSE][1] == "excel"


@requiere_postgres
def test_no_se_mueve_a_quien_ya_esta_en_otra_ficha(bd, api, tmp_path, ocr_simulado):
    db = bd.SessionLocal()
    ana = cabecera(crear_usuario(db, "ana@example.com"))
    db.close()
    otra = _subir(api, ana, tmp_path, ocr_simulado, "3590737.pdf", [CEDULA_AMARILLA])
    excel = _reporte(tmp_path / "reporte.xlsx", REPORTE)

    _subir(api, ana, tmp_path, ocr_simulado, "3640522.pdf", [CEDULA_DIGITAL], excel=excel)
    personas = _personas(bd)

    assert personas[ADRIANA][0] == uuid.UUID(otra["id"])                # se queda en su ficha
    assert personas[ADRIANA][1] != "excel"
    assert personas[JOSE][1] == "excel"


CEDULA_JOSE = (CEDULA_AMARILLA.replace("1.006.459.309", "1.119.580.110")
               .replace("ZAPATA CUELLAR", "LOAIZA ICO").replace("ADRIANA ALEJANDRA", "JOSE RUBIEL"))


@requiere_postgres
def test_al_reemplazar_quien_falta_en_pdf_no_toma_el_pdf_anterior(bd, api, tmp_path, ocr_simulado):
    db = bd.SessionLocal()
    ana = cabecera(crear_usuario(db, "ana@example.com"))
    db.close()
    excel = _reporte(tmp_path / "reporte.xlsx", REPORTE)

    vieja = _subir(api, ana, tmp_path, ocr_simulado, "3640522.pdf", [CEDULA_AMARILLA, CEDULA_DIGITAL], excel=excel)
    nueva = _subir(api, ana, tmp_path, ocr_simulado, "3640522 v2.pdf", [CEDULA_AMARILLA, CEDULA_DIGITAL], excel=excel,
                   acciones=[{"archivo": "3640522 v2.pdf", "accion": "reemplazar", "documento_id": vieja["id"]}])

    db = bd.SessionLocal()
    jose = db.query(Persona).filter(Persona.numero_identificacion == JOSE).one()
    db.close()
    assert jose.documento_id == uuid.UUID(nueva["id"])
    assert jose.documento_pdf_id is None and jose.motor_ocr == "excel"
    r = api.get("/api/personas", headers=ana, params={"documento_id": nueva["id"]})
    por_id = {p["numero_identificacion"]: p for p in r.json()}
    assert set(por_id) == {ADRIANA, ANYI, JOSE}
    assert por_id[JOSE]["en_pdf"] is False


@requiere_postgres
def test_si_luego_aparece_en_un_pdf_deja_de_faltar(bd, api, tmp_path, ocr_simulado):
    db = bd.SessionLocal()
    ana = cabecera(crear_usuario(db, "ana@example.com"))
    db.close()
    excel = _reporte(tmp_path / "reporte.xlsx", REPORTE)

    ficha = _subir(api, ana, tmp_path, ocr_simulado, "3640522.pdf", [CEDULA_AMARILLA, CEDULA_DIGITAL], excel=excel)
    parte = _subir(api, ana, tmp_path, ocr_simulado, "3640522 parte 2.pdf", [CEDULA_JOSE],
                   acciones=[{"archivo": "3640522 parte 2.pdf", "accion": "agregar", "documento_id": ficha["id"]}])

    r = api.get("/api/personas", headers=ana, params={"documento_id": ficha["id"]})
    por_id = {p["numero_identificacion"]: p for p in r.json()}
    assert set(por_id) == {ADRIANA, ANYI, JOSE}
    assert por_id[JOSE]["en_pdf"] is True
    assert por_id[JOSE]["documento_pdf_id"] == parte["id"]
    assert por_id[JOSE]["pagina_frente"] == 1
