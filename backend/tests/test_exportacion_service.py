"""
Pruebas para el Servicio de Exportación a Excel:
Verificación de columna 'Tipo de Documento' y columna final 'Estado / Hallazgo'
(Válido, No está en el PDF, No está en el Excel).
"""
import sys
import unittest
from unittest.mock import MagicMock, patch
from datetime import date, datetime
from openpyxl import load_workbook
from pathlib import Path

sys.path.insert(0, "backend")

from app.services.exportacion_service import exportacion_service


class TestExportacionService(unittest.TestCase):

    def test_columnas_exportacion(self):
        """Verifica que COLUMNAS tenga 9 columnas incluyendo 'Tipo de Documento' y 'Estado / Hallazgo'"""
        cols = list(exportacion_service.COLUMNAS.values())
        self.assertEqual(len(cols), 9)
        self.assertIn("Tipo de Documento", cols)
        self.assertIn("Estado / Hallazgo", cols)
        self.assertEqual(cols[1], "Tipo de Documento")
        self.assertEqual(cols[-1], "Estado / Hallazgo")

    def test_resolver_tipo_documento(self):
        """Verifica la resolución de tipos de documento a texto legible formal"""
        self.assertEqual(exportacion_service._resolver_tipo_documento("CEDULA_CIUDADANIA"), "Cédula de Ciudadanía (CC)")
        self.assertEqual(exportacion_service._resolver_tipo_documento("CC"), "Cédula de Ciudadanía (CC)")
        self.assertEqual(exportacion_service._resolver_tipo_documento("TARJETA_IDENTIDAD"), "Tarjeta de Identidad (TI)")
        self.assertEqual(exportacion_service._resolver_tipo_documento("TI"), "Tarjeta de Identidad (TI)")
        self.assertEqual(exportacion_service._resolver_tipo_documento("CEDULA_EXTRANJERIA"), "Cédula de Extranjería (CE)")
        self.assertEqual(exportacion_service._resolver_tipo_documento("PPT"), "Permiso por Protección Temporal (PPT)")
        self.assertEqual(exportacion_service._resolver_tipo_documento("CONTRASEÑA"), "Contraseña")
        self.assertEqual(exportacion_service._resolver_tipo_documento("PASAPORTE"), "Pasaporte")
        # Por texto OCR
        self.assertEqual(exportacion_service._resolver_tipo_documento(None, "REPUBLICA DE COLOMBIA TARJETA DE IDENTIDAD"), "Tarjeta de Identidad (TI)")
        self.assertEqual(exportacion_service._resolver_tipo_documento(None, "PERMISO POR PROTECCION TEMPORAL PPT"), "Permiso por Protección Temporal (PPT)")

    @patch("app.routers.personas._obtener_ids_en_excel")
    def test_exportar_personas_con_hallazgos_y_tipo(self, mock_obtener_ids):
        """
        Verifica que exportar_personas genere correctamente las columnas de Tipo de Documento y
        Estado / Hallazgo con 'Válido', 'No está en el PDF' y 'No está en el Excel'.
        """
        mock_obtener_ids.return_value = ({"1001", "1003"}, True)

        mock_db = MagicMock()
        mock_doc = MagicMock()
        mock_doc.nombre_original = "DOCUMENTOS_LOTE.pdf"

        # Persona 1: En PDF y en Excel -> Válido
        p1 = MagicMock()
        p1.id = "p-1"
        p1.numero_identificacion = "1001"
        p1.tipo_documento = "CEDULA_CIUDADANIA"
        p1.nombres = "JUAN"
        p1.apellidos = "PEREZ"
        p1.nombre_completo = "JUAN PEREZ"
        p1.fecha_nacimiento = date(1990, 5, 10)
        p1.edad = 34
        p1.documento = mock_doc
        p1.documento_id = "doc-1"
        p1.documento_pdf_id = None
        p1.requiere_revision = False
        p1.estado_registro = "VALID"
        p1.motor_ocr = "google_document_ai"
        p1.detalles_campos = {"en_pdf": True}
        p1.texto_ocr_crudo = "CEDULA DE CIUDADANIA 1001"
        p1.fecha_registro = datetime(2026, 9, 29, 10, 0)

        # Persona 2: En PDF pero NO en Excel -> No está en el Excel
        p2 = MagicMock()
        p2.id = "p-2"
        p2.numero_identificacion = "1002"
        p2.tipo_documento = "TARJETA_IDENTIDAD"
        p2.nombres = "CARLOS"
        p2.apellidos = "GOMEZ"
        p2.nombre_completo = "CARLOS GOMEZ"
        p2.fecha_nacimiento = date(2010, 1, 15)
        p2.edad = 14
        p2.documento = mock_doc
        p2.documento_id = "doc-1"
        p2.documento_pdf_id = None
        p2.requiere_revision = True
        p2.estado_registro = "REVIEW_REQUIRED"
        p2.motor_ocr = "google_document_ai"
        p2.detalles_campos = {"en_pdf": True}
        p2.texto_ocr_crudo = "TARJETA DE IDENTIDAD 1002"
        p2.fecha_registro = datetime(2026, 9, 29, 10, 5)

        # Persona 3: Solo en Excel (no en PDF) -> No está en el PDF
        p3 = MagicMock()
        p3.id = "p-3"
        p3.numero_identificacion = "1003"
        p3.tipo_documento = "CEDULA_CIUDADANIA"
        p3.nombres = "MARIA"
        p3.apellidos = "RODRIGUEZ"
        p3.nombre_completo = "MARIA RODRIGUEZ"
        p3.fecha_nacimiento = date(1985, 3, 20)
        p3.edad = 39
        p3.documento = None
        p3.documento_id = None
        p3.documento_pdf_id = None
        p3.requiere_revision = True
        p3.estado_registro = "REVIEW_REQUIRED"
        p3.motor_ocr = "excel"
        p3.detalles_campos = {"en_pdf": False, "origen": "excel_no_encontrado_en_pdf"}
        p3.texto_ocr_crudo = ""
        p3.fecha_registro = datetime(2026, 9, 29, 10, 10)

        # Persona 4: En PDF, pero aprobada manualmente por el usuario -> Válido
        p4 = MagicMock()
        p4.id = "p-4"
        p4.numero_identificacion = "1004"
        p4.tipo_documento = "PPT"
        p4.nombres = "ANA"
        p4.apellidos = "CASTRO"
        p4.nombre_completo = "ANA CASTRO"
        p4.fecha_nacimiento = date(1995, 7, 7)
        p4.edad = 29
        p4.documento = mock_doc
        p4.documento_id = "doc-1"
        p4.documento_pdf_id = None
        p4.requiere_revision = False
        p4.estado_registro = "VALID"
        p4.motor_ocr = "google_document_ai"
        p4.detalles_campos = {"aprobado_manual": True}
        p4.texto_ocr_crudo = "PPT 1004"
        p4.fecha_registro = datetime(2026, 9, 29, 10, 15)

        query_mock = MagicMock()
        query_mock.options.return_value = query_mock
        query_mock.outerjoin.return_value = query_mock
        query_mock.filter.return_value = query_mock
        query_mock.order_by.return_value = query_mock
        query_mock.all.return_value = [p1, p2, p3, p4]
        mock_db.query.return_value = query_mock

        ruta_generada = exportacion_service.exportar_personas(mock_db, {"usuario_id": "u-1"})
        self.assertTrue(Path(ruta_generada).exists())

        # Abrir el workbook y validar
        wb = load_workbook(ruta_generada)
        ws = wb["Personas Registradas"]

        # Encabezados en fila 3
        headers = [ws.cell(row=3, column=col).value for col in range(1, 10)]
        self.assertEqual(headers, [
            "Número Identificación",
            "Tipo de Documento",
            "Nombre Completo",
            "Fecha Nacimiento",
            "Edad",
            "Documento PDF Origen",
            "Requiere Revisión",
            "Fecha Registro",
            "Estado / Hallazgo"
        ])

        # Fila 4 (p1): Cédula de Ciudadanía, Válido
        self.assertEqual(ws.cell(row=4, column=1).value, "1001")
        self.assertEqual(ws.cell(row=4, column=2).value, "Cédula de Ciudadanía (CC)")
        self.assertEqual(ws.cell(row=4, column=9).value, "Válido")

        # Fila 5 (p2): Tarjeta de Identidad, No está en el Excel
        self.assertEqual(ws.cell(row=5, column=1).value, "1002")
        self.assertEqual(ws.cell(row=5, column=2).value, "Tarjeta de Identidad (TI)")
        self.assertEqual(ws.cell(row=5, column=9).value, "No está en el Excel")

        # Fila 6 (p3): Cédula de Ciudadanía, No está en el PDF
        self.assertEqual(ws.cell(row=6, column=1).value, "1003")
        self.assertEqual(ws.cell(row=6, column=2).value, "Cédula de Ciudadanía (CC)")
        self.assertEqual(ws.cell(row=6, column=9).value, "No está en el PDF")

        # Fila 7 (p4): PPT, Válido
        self.assertEqual(ws.cell(row=7, column=1).value, "1004")
        self.assertEqual(ws.cell(row=7, column=2).value, "Permiso por Protección Temporal (PPT)")
        self.assertEqual(ws.cell(row=7, column=9).value, "Válido")

        wb.close()
        # Limpieza
        try:
            Path(ruta_generada).unlink()
        except Exception:
            pass


if __name__ == "__main__":
    unittest.main()
