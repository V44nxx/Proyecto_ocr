"""
Regresiones corregidas en la versión 2.2.1.

Los tests sin base de datos corren siempre. Los que necesitan PostgreSQL real (migraciones del
arranque, índices, consultas por usuario) se ejecutan solo si TEST_DATABASE_URL apunta a una base
de pruebas desechable: cada test borra y recrea el esquema `public`.
"""
import os
import uuid
from datetime import date
from types import SimpleNamespace
from unittest.mock import MagicMock

import openpyxl
import pandas as pd
import pytest
from fastapi import HTTPException

from app.config import settings
from app.models.comparacion import Comparacion
from app.models.diferencia import Diferencia
from app.models.documento import Documento
from app.models.persona import Persona
from app.models.usuario import Usuario
from app.routers.comparacion import CorregirCampoRequest, corregir_campo_desde_comparacion
from app.routers.documentos import CancelarSubidaRequest, cancelar_subida_o_procesamiento
from app.services.google_document_ai_service import GoogleDocumentAIService


# ── Sin base de datos ───────────────────────────────────────────────

@pytest.mark.parametrize("texto, esperado", [
    ("NUIP1098765432", "NUIP1098765432"),
    ("CC1012345678", "CC1012345678"),
    ("No1.098.765", "No1.098.765"),
    ("MART1NEZ", "MARTINEZ"),
    ("PEREZ1", "PEREZI"),
    ("SAB!", "SABI"),
])
def test_normalizar_docai_no_rompe_cedulas_pegadas(texto, esperado):
    assert GoogleDocumentAIService.normalizar_texto_docai(texto) == esperado


def _db_corregir(persona):
    comparacion = SimpleNamespace(id=uuid.uuid4(), total_registros_excel=1, total_faltantes_bd=0,
                                  total_diferentes=0, total_coincidentes=0)
    db = MagicMock()

    def query(modelo):
        q = MagicMock()
        if modelo is Comparacion:
            q.filter.return_value.first.return_value = comparacion
        elif modelo is Persona:
            q.filter.return_value.first.return_value = persona
        else:
            q.filter.return_value.first.return_value = None
            q.filter.return_value.distinct.return_value.count.return_value = 0
        return q

    db.query.side_effect = query
    return db


def test_corregir_fecha_desde_comparacion_guarda_la_fecha():
    persona = Persona(id=uuid.uuid4(), numero_identificacion="1000000001", campos_revisados=[])
    usuario = SimpleNamespace(id=uuid.uuid4())
    res = corregir_campo_desde_comparacion(
        str(uuid.uuid4()),
        CorregirCampoRequest(numero_identificacion="1000000001", campo="fecha_nacimiento", nuevo_valor="08/05/1987"),
        db=_db_corregir(persona), usuario=usuario,
    )
    assert persona.fecha_nacimiento == date(1987, 5, 8)
    assert res["nuevo_valor"] == "1987-05-08"


def test_corregir_fecha_invalida_responde_400():
    persona = Persona(id=uuid.uuid4(), numero_identificacion="1000000001", campos_revisados=[])
    with pytest.raises(HTTPException) as exc:
        corregir_campo_desde_comparacion(
            str(uuid.uuid4()),
            CorregirCampoRequest(numero_identificacion="1000000001", campo="fecha_expedicion", nuevo_valor="no es fecha"),
            db=_db_corregir(persona), usuario=SimpleNamespace(id=uuid.uuid4()),
        )
    assert exc.value.status_code == 400


def test_corregir_nombre_completo_deja_nombres_coherentes():
    persona = Persona(id=uuid.uuid4(), numero_identificacion="1000000001", nombres="LUIS",
                      apellidos="RODRIGUEZ", nombre_completo="LUIS RODRIGUEZ", campos_revisados=[])
    corregir_campo_desde_comparacion(
        str(uuid.uuid4()),
        CorregirCampoRequest(numero_identificacion="1000000001", campo="nombre_completo", nuevo_valor="LUIS PEREZ"),
        db=_db_corregir(persona), usuario=SimpleNamespace(id=uuid.uuid4()),
    )
    assert (persona.nombre_completo, persona.nombres, persona.apellidos) == ("LUIS PEREZ", "LUIS PEREZ", None)
    assert "nombre_completo" in persona.campos_revisados


def test_cancelar_ids_inexistentes_no_cancela_otros_procesos():
    usuario = Usuario(id=uuid.uuid4(), email="a@example.com", rol="usuario")
    otro_en_proceso = Documento(id=uuid.uuid4(), usuario_id=usuario.id, nombre_original="otra.pdf",
                                nombre_archivo="otra.pdf", estado="procesando")
    db = MagicMock()
    q = MagicMock()
    # 1.ª consulta: los IDs pedidos (ya no existen). 2.ª: todo lo que está en proceso.
    q.filter.return_value.all.side_effect = [[], [otro_en_proceso]]
    db.query.return_value = q

    res = cancelar_subida_o_procesamiento(
        payload=CancelarSubidaRequest(documento_ids=[str(uuid.uuid4())]), db=db, usuario=usuario
    )

    assert res["eliminados"] == 0
    db.delete.assert_not_called()


def test_exportar_solo_menores(tmp_path, monkeypatch):
    from app.routers import personas as personas_router
    from app.services.exportacion_service import exportacion_service

    monkeypatch.setattr(settings, "EXPORT_DIR", str(tmp_path))
    monkeypatch.setattr(personas_router, "_obtener_ids_en_excel", lambda db, uid: (set(), False))

    def persona(cedula, nacimiento):
        return SimpleNamespace(
            id=uuid.uuid4(), numero_identificacion=cedula, nombres="ANA", apellidos="DIAZ", nombre_completo="ANA DIAZ",
            fecha_nacimiento=nacimiento, edad=Persona.edad.fget(SimpleNamespace(fecha_nacimiento=nacimiento)),
            tipo_documento="TARJETA_IDENTIDAD", texto_ocr_crudo="", documento=None, documento_id=None,
            documento_pdf_id=None, detalles_campos={}, motor_ocr="google_document_ai", requiere_revision=False,
            estado_registro="VALID", fecha_registro=None,
        )

    hoy = date.today()
    menor = persona("1000000010", date(hoy.year - 10, 1, 1))
    mayor = persona("1000000020", date(hoy.year - 30, 1, 1))
    q = MagicMock()
    for metodo in ("options", "outerjoin", "filter", "order_by"):
        getattr(q, metodo).return_value = q
    q.all.return_value = [menor, mayor]
    db = MagicMock()
    db.query.return_value = q

    ruta = exportacion_service.exportar_personas(db, {"usuario_id": str(uuid.uuid4()), "solo_menores": True})

    ws = openpyxl.load_workbook(ruta).active
    cedulas = {str(c.value) for row in ws.iter_rows() for c in row}
    assert "1000000010" in cedulas
    assert "1000000020" not in cedulas


# ── Con PostgreSQL real ─────────────────────────────────────────────

TEST_DATABASE_URL = os.getenv("TEST_DATABASE_URL")
requiere_postgres = pytest.mark.skipif(not TEST_DATABASE_URL, reason="Define TEST_DATABASE_URL para correr estos tests")


@pytest.fixture
def bd(monkeypatch, tmp_path):
    from sqlalchemy import create_engine, text
    from sqlalchemy.orm import sessionmaker
    import app.database as database

    engine = create_engine(TEST_DATABASE_URL)
    with engine.begin() as conn:
        # Cerrar sesiones que un test fallido haya dejado abiertas, para que el DROP no se bloquee
        conn.execute(text(
            "SELECT pg_terminate_backend(pid) FROM pg_stat_activity "
            "WHERE datname = current_database() AND pid <> pg_backend_pid()"
        ))
        conn.execute(text("DROP SCHEMA public CASCADE; CREATE SCHEMA public;"))
    monkeypatch.setattr(database, "engine", engine)
    monkeypatch.setattr(database, "SessionLocal", sessionmaker(autocommit=False, autoflush=False, bind=engine))
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setattr(settings, "EXPORT_DIR", str(tmp_path / "exports"))
    monkeypatch.setattr(settings, "ADMIN_EMAIL", None)
    database.create_tables()
    yield database
    engine.dispose()


def _usuario(db, email, rol="usuario"):
    u = Usuario(email=email, nombre=email, password_hash="x", rol=rol)
    db.add(u)
    db.commit()
    return u


def _ficha(db, usuario, nombre, paginas, **extra):
    d = Documento(usuario_id=usuario.id, nombre_archivo=extra.pop("nombre_archivo", f"{uuid.uuid4()}_{nombre}"),
                  nombre_original=nombre, total_paginas=paginas, estado=extra.pop("estado", "completado"), **extra)
    db.add(d)
    db.commit()
    return d


def _persona(db, usuario, ficha, cedula, nombres, apellidos, **extra):
    p = Persona(usuario_id=usuario.id, documento_id=ficha.id if ficha else None, numero_identificacion=cedula,
                nombres=nombres, apellidos=apellidos, nombre_completo=extra.pop("nombre_completo", f"{nombres} {apellidos}"),
                motor_ocr=extra.pop("motor_ocr", "google_document_ai"), estado_registro="VALID",
                requiere_revision=False, confianza_extraccion=95, grupo_documento_id="G-1", **extra)
    db.add(p)
    db.commit()
    return p


@requiere_postgres
def test_arranque_no_mueve_personas_de_fichas_pequenas(bd):
    db = bd.SessionLocal()
    u = _usuario(db, "a@example.com")
    grande = _ficha(db, u, "3000001.pdf", 10)
    pequena = _ficha(db, u, "3000002.pdf", 2)
    p = _persona(db, u, pequena, "1000000001", "ANA", "TORRES")
    ids = (p.id, pequena.id)
    db.close()

    bd.create_tables()

    db = bd.SessionLocal()
    assert db.get(Persona, ids[0]).documento_id == ids[1]
    assert db.get(Documento, ids[1]).visible_en_subida is True
    db.close()


@requiere_postgres
def test_arranque_devuelve_pdf_individual_a_su_ficha(bd):
    db = bd.SessionLocal()
    u = _usuario(db, "a@example.com")
    lote = _ficha(db, u, "3000001.pdf", 10)
    individual = _ficha(db, u, "cedula.pdf", 1, nombre_archivo="cedula_1000000001_ab12cd34.pdf")
    p = _persona(db, u, individual, "1000000001", "ANA", "TORRES")
    lote.metadatos = {"personas_extraidas_datos": [{"id": str(p.id)}]}
    db.commit()
    ids = (p.id, lote.id, individual.id)
    db.close()

    bd.create_tables()

    db = bd.SessionLocal()
    persona = db.get(Persona, ids[0])
    assert (persona.documento_id, persona.documento_pdf_id) == (ids[1], ids[2])
    assert db.get(Documento, ids[2]).visible_en_subida is False
    db.close()


@requiere_postgres
def test_arranque_respeta_nombre_corregido_a_mano(bd):
    db = bd.SessionLocal()
    u = _usuario(db, "a@example.com")
    f = _ficha(db, u, "3000001.pdf", 10)
    p = _persona(db, u, f, "1000000002", "LUIS", "RODRIGUEZ", nombre_completo="LUIS PEREZ",
                 campos_revisados=["nombre_completo"])
    pid = p.id
    db.close()

    bd.create_tables()

    db = bd.SessionLocal()
    assert db.get(Persona, pid).nombre_completo == "LUIS PEREZ"
    db.close()


@requiere_postgres
def test_reporte_de_comparacion_solo_trae_personas_del_dueno(bd, tmp_path):
    from app.services.comparacion_service import comparacion_service

    db = bd.SessionLocal()
    a = _usuario(db, "a@example.com")
    b = _usuario(db, "b@example.com")
    _persona(db, a, _ficha(db, a, "3000001.pdf", 4), "1000000002", "LUIS", "RODRIGUEZ")
    _persona(db, b, _ficha(db, b, "3000009.pdf", 4), "2000000009", "PERSONA", "AJENA")
    excel = tmp_path / "planilla.xlsx"
    pd.DataFrame([{"Cedula": "1000000002", "Nombre completo": "LUIS RODRIGUEZ"}]).to_excel(excel, index=False)
    comp = Comparacion(usuario_id=a.id, nombre_archivo="planilla.xlsx", nombre_original="planilla.xlsx",
                       ruta_archivo=str(excel), estado="pendiente")
    db.add(comp)
    db.commit()

    comparacion_service.ejecutar_comparacion(str(comp.id), str(excel), db)
    ruta = comparacion_service.generar_reporte_xlsx(str(comp.id), db)

    celdas = {str(c.value) for ws in openpyxl.load_workbook(ruta).worksheets for row in ws.iter_rows() for c in row}
    assert "1000000002" in celdas
    assert "2000000009" not in celdas
    db.close()


@requiere_postgres
def test_reconciliacion_por_nombre_corrige_la_cedula(bd, tmp_path):
    from app.services.comparacion_service import comparacion_service

    db = bd.SessionLocal()
    u = _usuario(db, "c@example.com")
    leida = _persona(db, u, _ficha(db, u, "3000010.pdf", 4), "1000000077", "MARIA", "LOPEZ GOMEZ",
                     fecha_nacimiento=date(1990, 1, 1))
    # Registro creado desde la planilla, sin PDF, con la cédula correcta
    _persona(db, u, None, "1000000078", "MARIA", "LOPEZ GOMEZ", motor_ocr="manual",
             detalles_campos={"en_pdf": False}, lugar_expedicion="FLORENCIA")
    excel = tmp_path / "planilla.xlsx"
    pd.DataFrame([{"Cedula": "1000000078", "Nombres": "MARIA", "Apellidos": "LOPEZ GOMEZ"}]).to_excel(excel, index=False)
    comp = Comparacion(usuario_id=u.id, nombre_archivo="p.xlsx", nombre_original="p.xlsx", ruta_archivo=str(excel),
                       estado="pendiente")
    db.add(comp)
    db.commit()
    ids = (leida.id, comp.id, u.id)

    comparacion_service.ejecutar_comparacion(str(comp.id), str(excel), db)
    db.close()

    db = bd.SessionLocal()
    assert db.get(Comparacion, ids[1]).estado == "completado"
    personas = db.query(Persona).filter(Persona.usuario_id == ids[2]).all()
    assert len(personas) == 1
    assert personas[0].id == ids[0]
    assert personas[0].numero_identificacion == "1000000078"
    assert personas[0].fecha_nacimiento == date(1990, 1, 1)
    assert personas[0].lugar_expedicion == "FLORENCIA"
    db.close()


def _datos_ocr(cedula, nombres, apellidos, motivos=None):
    return {"identificacion": cedula, "nombres": nombres, "apellidos": apellidos, "confianza_extraccion": 40.0,
            "pagina_frente": 1, "grupo_documento_id": f"G-{nombres}",
            "detalles_campos": {"motivos_revision": motivos} if motivos else {}}


@requiere_postgres
def test_unificar_persona_con_motivos_de_revision(bd):
    from app.services.ocr_service import ocr_service

    db = bd.SessionLocal()
    u = _usuario(db, "d@example.com")
    f1 = _ficha(db, u, "3000020.pdf", 4)
    f2 = _ficha(db, u, "3000021.pdf", 4)
    assert ocr_service._guardar_persona(_datos_ocr("1098765432", "CARLOS", None, ["Falta apellido"]), "Frente Pag 1", str(f1.id), db)
    segunda = ocr_service._guardar_persona(_datos_ocr("1098765432", "CARLOS", None, ["Falta apellido"]), "Frente Pag 1", str(f2.id), db)
    assert segunda is not None
    assert db.query(Persona).filter(Persona.usuario_id == u.id).one().documento_id == f2.id
    db.close()


@requiere_postgres
def test_varias_personas_sin_cedula_en_la_misma_cuenta(bd):
    from app.services.ocr_service import ocr_service

    db = bd.SessionLocal()
    u = _usuario(db, "e@example.com")
    f = _ficha(db, u, "3000030.pdf", 4)
    guardadas = [ocr_service._guardar_persona(_datos_ocr(None, n, "SUAREZ"), "Frente Pag 1", str(f.id), db)
                 for n in ("PEDRO", "PABLO")]
    assert all(guardadas)
    assert db.query(Persona).filter(Persona.usuario_id == u.id).count() == 2
    db.close()


@requiere_postgres
def test_eliminar_usuario_borra_sus_datos_y_no_pasan_al_admin(bd, tmp_path):
    from app.routers.auth import eliminar_usuario

    db = bd.SessionLocal()
    admin = _usuario(db, "admin@example.com", rol="admin")
    x = _usuario(db, "x@example.com")
    pdf = tmp_path / "ficha.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    f = _ficha(db, x, "3000040.pdf", 6, ruta_archivo=str(pdf))
    _persona(db, x, f, "1000000500", "ROSA", "DIAZ")
    db.add(Comparacion(usuario_id=x.id, nombre_archivo="p.xlsx", nombre_original="p.xlsx", estado="completado"))
    db.commit()

    res = eliminar_usuario(str(x.id), usuario_actual=admin, db=db)
    db.close()
    bd.create_tables()

    db = bd.SessionLocal()
    assert (res["fichas_eliminadas"], res["personas_eliminadas"]) == (1, 1)
    assert db.query(Persona).count() == 0
    assert db.query(Documento).count() == 0
    assert db.query(Comparacion).count() == 0
    assert not pdf.exists()
    db.close()


@requiere_postgres
def test_cancelar_no_borra_una_ficha_completada(bd):
    db = bd.SessionLocal()
    u = _usuario(db, "f@example.com")
    f = _ficha(db, u, "3000050.pdf", 4, estado="completado")
    _persona(db, u, f, "1000000600", "LUZ", "MORA")

    res = cancelar_subida_o_procesamiento(
        payload=CancelarSubidaRequest(documento_ids=[str(f.id)]), db=db, usuario=u
    )

    assert res["eliminados"] == 0
    assert db.query(Persona).filter(Persona.documento_id == f.id).count() == 1
    db.close()
